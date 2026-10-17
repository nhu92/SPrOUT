"""Local-only export and truth evaluation after ten TTU pipelines finish.

Never import this module in the inference payload. Workbook cells are data.
"""
import argparse
from collections import Counter
import csv
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import shutil
from zipfile import ZipFile, ZIP_DEFLATED


METHODS = ('merged', 'c_style', 'balanced_z', 'standalone_ml', 'core', 'auxiliary', 'review')
TRUTH_SHA256 = '421cca29377ab01e57cb6aeded7bf446ba4faf06282120a4ed602e8c67bbb01f'


def read(path):
    return json.loads(Path(path).read_text(encoding='utf-8'))


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def dump(path, value):
    Path(path).write_text(json.dumps(value, indent=2, allow_nan=False)+'\n', encoding='utf-8')


def names(path):
    lines = Path(path).read_text(encoding='utf-8').splitlines()
    if len(lines) != len(set(lines)) or any(not x.strip() for x in lines):
        raise ValueError('Duplicate or empty genus in '+str(path))
    return set(lines)


def write_csv(path, rows, fields=None):
    fields = fields or list(dict.fromkeys(k for r in rows for k in r))
    with Path(path).open('w', newline='', encoding='utf-8-sig') as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def read_csv(path):
    with Path(path).open(newline='', encoding='utf-8-sig') as f:
        return list(csv.DictReader(f))


def complete(batch):
    batch = Path(batch)
    audit = read(batch/'AUDIT.json')
    if not (batch/'PREDICTIONS_COMPLETE').exists() or not audit.get('all_10_complete') or not audit.get('code_model_inputs_products_verified'):
        raise ValueError('All ten pipelines and immutable-input audits must finish first')
    selection = read(batch/'selection.json')['selected']
    calls = read(batch/'predictions.json')
    if len(selection) != 10 or len(calls) != 10:
        raise ValueError('Expected exactly ten mixtures')
    lookup = {r['sample']:r for r in calls}
    if len(lookup) != 10 or set(lookup) != {r['sample'] for r in selection}:
        raise ValueError('Cohort mismatch')
    for row in selection:
        path = batch/'samples'/row['project']
        if not (path/'SUCCESS').exists() or (path/'FAILED').exists():
            raise ValueError('Incomplete mixture: '+row['sample'])
        if read(path/'predictions.json') != lookup[row['sample']]:
            raise ValueError('Combined prediction mismatch')
        for method in METHODS:
            if names(path/(method+'_genera.txt')) != set(lookup[row['sample']][method]):
                raise ValueError('Text/JSON prediction mismatch')
    return selection, lookup


def export(batch, output):
    """Seal all results before a truth-reading function can run."""
    batch, output = Path(batch), Path(output)
    selection, calls = complete(batch)
    if output.exists():
        raise FileExistsError('Use a new prediction export directory')
    output.mkdir(parents=True)
    (output/'.gitattributes').write_text('* -text whitespace=cr-at-eol\n', encoding='utf-8')
    for name in ('PROTOCOL.json', 'selection.json', 'predictions.json', 'SEARCHES.json', 'AUDIT.json', 'PREDICTIONS_COMPLETE'):
        shutil.copy2(batch/name, output/name)
    predictions, evidence, search_regions = [], [], []
    with ZipFile(output/'prediction_sets.zip', 'w', compression=ZIP_DEFLATED) as archive:
        for row in selection:
            sample, project = row['sample'], row['project']
            path = batch/'samples'/project
            table = read_csv(path/'genus_evidence.csv') if (path/'genus_evidence.csv').exists() else []
            by_genus = {r['genus']:r for r in table}
            for r in table:
                evidence.append({'sample':sample, **r})
            search = read(path/'SEARCH.json')
            for method in METHODS:
                archive.write(path/(method+'_genera.txt'), sample+'/'+method+'_genera.txt')
                for genus in sorted(calls[sample][method]):
                    r = by_genus[genus]
                    if r[method] != 'True':
                        raise ValueError('Evidence/selected prediction mismatch')
                    predictions.append({'sample':sample, 'genus':genus, 'family':r['family'], 'method':method,
                                        'ml_gene_votes':r['ml_gene_votes'], 'nearest_gene_votes':r['nearest_gene_votes'],
                                        'search_status':search['search_status']})
            for stage, file in [('initial', 'initial_genera.txt'), ('final_search', 'final_search_genera.txt'), ('unresolved', 'unresolved_genera.txt')]:
                search_regions.extend({'sample':sample, 'stage':stage, 'genus':g} for g in sorted(names(path/file)))
            for directory in sorted(path.glob('round_*')):
                for stage, file in [('before', 'candidate_genera.txt'), ('after', 'retained_genera.txt')]:
                    search_regions.extend({'sample':sample, 'stage':directory.name+'_'+stage, 'genus':g}
                                          for g in sorted(names(directory/file)))
    fields = ['sample','genus','family','method','ml_gene_votes','nearest_gene_votes','search_status']
    write_csv(output/'all_method_predictions.csv', predictions, fields)
    write_csv(output/'merged_predictions.csv', [r for r in predictions if r['method']=='merged'], fields)
    write_csv(output/'candidate_evidence.csv', evidence)
    write_csv(output/'search_regions.csv', search_regions)
    record = {'rank':'genus', 'primary':'merged', 'samples':10, 'truth_read_by_exporter':False,
              'sealed_utc':datetime.now(timezone.utc).isoformat(),
              'evaluation_plan_sha256':sha(Path(__file__).with_name('EVALUATION_PLAN.md')),
              'artifact_sha256':{p.name:sha(p) for p in sorted(output.iterdir()) if p.is_file()}}
    dump(output/'PREDICTION_EXPORT.json', record)
    print(json.dumps({'sealed_samples':10, 'primary_genus_events':sum(r['method']=='merged' for r in predictions)}))


def verify_export(output):
    record = read(Path(output)/'PREDICTION_EXPORT.json')
    if record['evaluation_plan_sha256'] != sha(Path(__file__).with_name('EVALUATION_PLAN.md')):
        raise ValueError('Evaluation plan changed after sealing')
    for name, digest in record['artifact_sha256'].items():
        if sha(Path(output)/name) != digest:
            raise ValueError('Sealed prediction artifact changed: '+name)


def table_truth(path, genus_column, species_name=False):
    """Read the literal table; a species column can supply its first-word genus."""
    from openpyxl import load_workbook
    from experiments.reference_atlas.real_truth import xml_sheet
    if sha(path) != TRUTH_SHA256:
        raise ValueError('Unexpected original 30-source workbook')
    book = load_workbook(path, read_only=True, data_only=True)
    sheet = book['Table S1']
    xml = xml_sheet(path, 'Table S1')
    truth, cells = {}, {}
    for row in range(4, sheet.max_row+1):
        identity = sheet[f'A{row}'].value
        if not isinstance(identity, int) or isinstance(identity, bool) or not 1 <= identity <= 30:
            continue
        addresses = [f'A{row}', f'{genus_column}{row}', f'H{row}']
        values = [sheet[a].value for a in addresses]
        for address, value in zip(addresses, values):
            if str(value).strip() != str(xml.get(address, '')).strip():
                raise ValueError('Independent truth-cell mismatch: '+address)
            cells[address] = value
        name = str(values[1]).strip()
        genus = name.split()[0] if species_name else name
        if not genus.isalpha() or not genus[0].isupper() or identity in truth:
            raise ValueError('Ambiguous or duplicate literal genus: '+repr(values))
        truth[identity] = {'source_id':identity, 'literal_name':name, 'genus':genus,
                           'family':str(values[2]).strip(), 'cells':addresses}
    book.close()
    if set(truth) != set(range(1,31)):
        raise ValueError('Expected all thirty original source identities')
    return truth, cells


def measures(expected, called):
    tp, fp, fn = len(expected & called), len(called-expected), len(expected-called)
    return counts(tp, fp, fn)


def counts(tp, fp, fn):
    return {'TP':tp, 'FP':fp, 'FN':fn, 'precision':tp/(tp+fp) if tp+fp else None,
            'recall':tp/(tp+fn) if tp+fn else None, 'F1':2*tp/(2*tp+fp+fn) if 2*tp+fp+fn else None}


def loss_stage(genus, family, row, path, called):
    if genus not in names(path/'initial_genera.txt'):
        if family not in row['families']:
            return 'absent_literal_parent_family'
        return 'absent_recoverable_initial_reference'
    for directory in sorted(path.glob('round_*')):
        if genus not in names(directory/'retained_genera.txt'):
            return directory.name+'_routing_exclusion'
    if genus in names(path/'unresolved_genera.txt'):
        return 'unresolved_search_region'
    if genus not in names(path/'final_search_genera.txt'):
        return 'outside_final_joint_search'
    panel = read(path/'final_reference/panel.json') if (path/'final_reference/panel.json').exists() else {'selection':[]}
    if genus not in {r['genus'] for r in panel['selection']}:
        return 'final_reference_QC_unavailable'
    if genus not in called:
        return 'final_classifier_rejection'
    return 'recovered'


def evaluate(batch, output, workbook, genus_column, species_name=False):
    batch, output = Path(batch), Path(output)
    selection, calls = complete(batch)
    verify_export(output)
    for name in ('PROTOCOL.json','selection.json','predictions.json','SEARCHES.json','AUDIT.json'):
        if sha(batch/name)!=sha(output/name):
            raise ValueError('Sealed batch differs from evaluation input: '+name)
    # The first workbook-content read occurs only after these completion checks.
    truth, cells = table_truth(workbook, genus_column, species_name)
    dump(output/'truth.json', {'view':'table_literal', 'workbook_sha256':sha(workbook), 'sheet':'Table S1',
                             'genus_column':genus_column, 'first_word_of_species_name':species_name,
                             'sources':list(truth.values()), 'independently_verified_used_cells':cells})
    by_sample, survival, losses, timings, rounds, phases = [], [], [], [], [], []
    for row in selection:
        sample = row['sample'];path = batch/'samples'/row['project']
        expected = {truth[int(i)]['genus'] for i in sample.split('x')}
        for method in METHODS:
            called = set(calls[sample][method])
            by_sample.append({'sample':sample, 'components':len(sample.split('x')), 'method':method,
                              **measures(expected,called), 'expected':';'.join(sorted(expected)),
                              'predicted':';'.join(sorted(called)), 'false_positives':';'.join(sorted(called-expected)),
                              'false_negatives':';'.join(sorted(expected-called))})
        # One loss event per genus, even if multiple source species share it.
        for genus in sorted(expected):
            sources = [truth[int(i)] for i in sample.split('x') if truth[int(i)]['genus']==genus]
            family_set = {s['family'] for s in sources}
            if len(family_set)!=1:
                raise ValueError('Conflicting literal genus parents')
            family = next(iter(family_set))
            losses.append({'sample':sample, 'genus':genus, 'family':family,
                           'literal_parent_in_search':family in row['families'],
                           'genus_in_initial_search':genus in names(path/'initial_genera.txt'),
                           'first_loss':loss_stage(genus, family, row, path, set(calls[sample]['merged']))})
        stages = [('initial', names(path/'initial_genera.txt'))]
        for directory in sorted(path.glob('round_*')):
            stages.append((directory.name, names(directory/'retained_genera.txt')))
        stages += [('final_search',names(path/'final_search_genera.txt')),
                   ('final_reference', {r['genus'] for r in read(path/'final_reference/panel.json')['selection']}
                    if (path/'final_reference/panel.json').exists() else set()),
                   ('merged',set(calls[sample]['merged']))]
        for stage, pool in stages:
            survival.append({'sample':sample, 'stage':stage, 'pool_genera':len(pool), 'expected_genera':len(expected),
                             'true_retained':len(expected&pool), 'missing_true_genera':';'.join(sorted(expected-pool))})
        search = read(path/'SEARCH.json')
        timings.append({'sample':sample, 'components':len(sample.split('x')), 'parent_families':len(row['families']),
                        'search_status':search['search_status'], 'rounds':len(search['rounds']),
                        **{k:search[k] for k in ('initial_genera','remaining_search_genera','final_search_genera','unresolved_genera',
                                                'seconds','new_download_bytes','fetch_seconds','qc_seconds','reference_alignment_seconds','checked_recoveries')},
                        'final_reference_genera':search['final'].get('reference_genera',0),
                        'final_reference_species':search['final'].get('reference_species',0),
                        'merged_genera':len(calls[sample]['merged'])})
        for r in search['rounds']:
            rounds.append({'sample':sample, **r})
        parts={'fetch_parse':search['fetch_seconds'],'reference_QC':search['qc_seconds'],
               'reference_projection':search['reference_alignment_seconds'],
               'query_align_tree':sum(r['tree_seconds'] for r in search['rounds'])+search['final'].get('tree_seconds',0)}
        other=search['seconds']-sum(parts.values())
        if other < -.01:
            raise ValueError('Phase timings exceed total search elapsed time')
        parts['planning_routing_scoring_other']=max(0,other)
        phases.extend({'sample':sample,'phase':phase,'seconds':seconds,'fraction_of_search':seconds/search['seconds']}
                      for phase,seconds in parts.items())
    summary=[]
    for method in METHODS:
        sub=[r for r in by_sample if r['method']==method]
        summary.append({'view':'table_literal','method':method,'samples':10,
                        **counts(*(sum(r[k] for r in sub) for k in ('TP','FP','FN'))),
                        'exact_sets':sum(r['FP']==r['FN']==0 for r in sub)})
    primary = next(r for r in summary if r['method']=='merged')
    if sum(r['first_loss']=='recovered' for r in losses)!=primary['TP'] or len(losses)!=primary['TP']+primary['FN']:
        raise ValueError('First-loss diagnosis must reproduce primary genus recall')
    complexity=[]
    for components in (3,6,10):
        for method in METHODS:
            sub=[r for r in by_sample if r['method']==method and r['components']==components]
            complexity.append({'components':components,'method':method,'samples':len(sub),
                               **counts(*(sum(r[k] for r in sub) for k in ('TP','FP','FN'))),
                               'exact_sets':sum(r['FP']==r['FN']==0 for r in sub)})
    for name, rows in [('summary.csv',summary),('by_sample.csv',by_sample),('truth_survival.csv',survival),
                       ('first_losses.csv',losses),('timings.csv',timings),('rounds.csv',rounds),
                       ('summary_by_complexity.csv',complexity),('runtime_phases.csv',phases)]:
        write_csv(output/name,rows)
    # Recount independently from plain files and the XML reader's source cells.
    from experiments.reference_atlas.real_truth import xml_sheet
    xml = xml_sheet(workbook,'Table S1')
    xml_truth = {i:(xml[t['cells'][1]].strip().split()[0] if species_name else xml[t['cells'][1]].strip()) for i,t in truth.items()}
    for result in by_sample:
        row = next(r for r in selection if r['sample']==result['sample'])
        expected={xml_truth[int(i)] for i in result['sample'].split('x')}
        called=names(batch/'samples'/row['project']/(result['method']+'_genera.txt'))
        recount=measures(expected,called)
        if any(result[k]!=recount[k] for k in ('TP','FP','FN')):
            raise ValueError('Independent candidate recount disagrees')
    if sha(workbook)!=TRUTH_SHA256:
        raise ValueError('Source workbook changed during read-only evaluation')
    dump(output/'EVALUATION_AUDIT.json',{'samples':10,'primary':'merged','view':'table_literal',
         'all_ten_predictions_sealed_before_truth_read':True,'openpyxl_XML_used_cells_agree':True,
         'plain_method_text_XML_recounts_agree':True,'workbook_sha256':sha(workbook),
         'source_workbook_modified':False,'first_losses':dict(Counter(r['first_loss'] for r in losses)),
         'evaluated_utc':datetime.now(timezone.utc).isoformat()})
    print(json.dumps(summary,indent=2))


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('command',choices=('export','evaluate'))
    parser.add_argument('--batch',required=True,type=Path)
    parser.add_argument('--output',required=True,type=Path)
    parser.add_argument('--workbook',type=Path)
    parser.add_argument('--genus-column')
    parser.add_argument('--species-name',action='store_true')
    args=parser.parse_args()
    if args.command=='export':
        export(args.batch,args.output)
    else:
        if not args.workbook or not args.genus_column:
            parser.error('evaluate requires --workbook and --genus-column')
        evaluate(args.batch,args.output,args.workbook,args.genus_column,args.species_name)
