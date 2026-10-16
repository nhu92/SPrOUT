"""Evaluate completed fusion artifacts, then independently audit exported ZIPs.

Prediction code has no truth argument. All input labels here are evaluated only
after the merged candidate lists have been produced and their hashes checked.
"""
import argparse
import csv
import gzip
import io
import json
from pathlib import Path
from zipfile import ZipFile

from .audit_transfer import xml_truth
from .evaluate_transfer import measures, write_csv
from .fuse import digest, decide, read_csv, write_json


TRUTH_HASH = "421cca29377ab01e57cb6aeded7bf446ba4faf06282120a4ed602e8c67bbb01f"
METHODS = ("baseline_c", "stable_c", "core", "merged", "standalone_ml", "expanded_review_union")
VIEWS = ("table_literal", "harmonized_two_taxa")
GROUPS = ("all100", "unique98", "historical30", "remaining70")


def called(sample, method):
    if method == "expanded_review_union":
        return set(sample["merged"]) | set(sample["review"])
    return set(sample[method])


def in_group(row, group):
    return (group == "all100" or group == "unique98" and row["first_composition"]
            or group == "historical30" and row["historical30"]
            or group == "remaining70" and not row["historical30"])


def evaluate(root, truth_path):
    root = Path(root)
    if not (root / "PREDICTIONS_COMPLETE").exists():
        raise ValueError("Predictions must complete before evaluation")
    manifest = json.loads((root / "MANIFEST.json").read_text())
    for name, expected in manifest["artifact_sha256"].items():
        if digest(root / name) != expected:
            raise ValueError(f"Prediction artifact changed: {name}")
    if digest(truth_path) != TRUTH_HASH:
        raise ValueError("Unexpected truth workbook")
    calls = json.loads((root / "predictions.json").read_text())
    policy = json.loads((root / "POLICY.json").read_text())
    evidence = list(csv.DictReader(io.StringIO(gzip.decompress((root / "evidence.csv.gz").read_bytes()).decode("utf-8-sig"))))
    universe = {r["family"] for r in evidence}
    if len(calls) != 100 or len(universe) != 362:
        raise ValueError("Unexpected sample/universe count")
    # Primary evaluator uses openpyxl; separate audit below uses ZIP/XML.
    from openpyxl import load_workbook
    workbook = load_workbook(truth_path, read_only=True, data_only=True)
    truth = {int(r[0]): str(r[7]).strip() for r in workbook["Table S1"].iter_rows(min_row=4, values_only=True) if isinstance(r[0], int)}
    workbook.close()
    if set(truth) != set(range(1, 31)):
        raise ValueError("Unexpected source identities")
    first, by_sample, effects = {}, [], []
    evidence_by_key = {(r["sample"], r["family"]): r for r in evidence}
    for sample in calls:
        composition = tuple(sorted(map(int, sample["sample"].split("x"))))
        first.setdefault(composition, sample["sample"])
        for view in VIEWS:
            correction = {1: "Cabombaceae", 25: "Capparaceae"} if view == "harmonized_two_taxa" else {}
            expected = {correction.get(i, truth[i]) for i in composition}
            for method in METHODS:
                predicted = called(sample, method)
                if not predicted <= universe:
                    raise ValueError("Prediction outside fixed universe")
                counts = map(len, (expected & predicted, predicted - expected, expected - predicted, universe - expected - predicted))
                by_sample.append({"sample": sample["sample"], "view": view, "method": method,
                    "historical30": sample["historical30"], "first_composition": first[composition] == sample["sample"],
                    **measures(*counts), "expected": ";".join(sorted(expected)), "predicted": ";".join(sorted(predicted)),
                    "false_positives": ";".join(sorted(predicted - expected)), "false_negatives": ";".join(sorted(expected - predicted))})
            core, merged, c = (set(sample[k]) for k in ("core", "merged", "baseline_c"))
            for event, families in (("auxiliary_added_to_core", merged - core), ("removed_from_C", c - merged),
                                    ("review_true_not_final", expected & set(sample["review"])),
                                    ("final_false_positive", merged - expected), ("final_false_negative", expected - merged)):
                for family in sorted(families):
                    e = evidence_by_key.get((sample["sample"], family), {})
                    effects.append({"sample": sample["sample"], "view": view, "event": event,
                        "family": family, "expected": family in expected, "ml_gene_votes": e.get("ml_gene_votes", ""),
                        "rescue_selected": e.get("rescue_selected", ""), "stability_lower": e.get("stability_lower", ""),
                        "discriminating_genes": e.get("discriminating_genes", ""), "comparable_genes": e.get("comparable_genes", "")})
    summary = []
    for group in GROUPS:
        for view in VIEWS:
            for method in METHODS:
                sub = [r for r in by_sample if r["view"] == view and r["method"] == method and in_group(r, group)]
                counts = [sum(r[k] for r in sub) for k in ("TP", "FP", "FN", "TN")]
                summary.append({"group": group, "view": view, "method": method, "samples": len(sub),
                    "negative_universe_families": len(universe), **measures(*counts),
                    "exact_sets": sum(r["FP"] == r["FN"] == 0 for r in sub)})
    write_csv(root / "summary.csv", summary)
    write_csv(root / "by_sample.csv", by_sample)
    write_csv(root / "error_and_route_details.csv", effects)
    audit(root, truth_path, calls, universe, evidence, policy, by_sample, summary)
    write_json(root / "EVALUATION.json", {"truth_sha256": digest(truth_path),
        "evaluator_sha256": digest(__file__), "policy_sha256": digest(root / "POLICY.json"),
        "prediction_manifest_sha256": digest(root / "MANIFEST.json"),
        "previous_sample_results_known_during_method_development": True,
        "independent_validation": False,
        "artifact_sha256": {name: digest(root / name) for name in
                            ("summary.csv", "by_sample.csv", "error_and_route_details.csv", "COUNT_AUDIT.json")}})
    print(json.dumps([r for r in summary if r["group"] == "all100" and r["method"] in {"core", "merged"}], indent=2))


def audit(root, truth_path, calls, universe, evidence, policy, by_sample, summary):
    """Recount original truth against exported per-sample lists, not CSV counts."""
    truth = xml_truth(truth_path)
    lookup = {r["sample"]: r for r in calls}
    zip_checks = 0
    with ZipFile(root / "prediction_sets.zip") as archive:
        if json.loads(archive.read("POLICY.json")) != policy:
            raise ValueError("ZIP policy mismatch")
        for sample in calls:
            for method in ("core", "auxiliary", "merged", "review"):
                exported = archive.read(f'{sample["project"]}/{method}_families.txt').decode().splitlines()
                if exported != sample[method]:
                    raise ValueError("Exported lists disagree with predictions")
                zip_checks += 1
            if set(sample["merged"]) != set(sample["core"]) | set(sample["auxiliary"]):
                raise ValueError("Route union mismatch")
            if set(sample["merged"]) & set(sample["review"]):
                raise ValueError("Review leaked into final")
    # Replay the declared rule against the frozen evidence snapshot.
    for r in evidence:
        simplified = {k: r[k] == "True" for k in ("c_selected", "rescue_selected", "b_selected")}
        simplified.update(ml_gene_votes=int(r["ml_gene_votes"]), stability_lower=float(r["stability_lower"]))
        result = decide(simplified, policy)
        for method in ("core", "auxiliary", "merged", "review"):
            if result[method] != (r["family"] in lookup[r["sample"]][method]):
                raise ValueError("Rule replay mismatch")
    for r in by_sample:
        correction = {1: "Cabombaceae", 25: "Capparaceae"} if r["view"] == "harmonized_two_taxa" else {}
        expected = {correction.get(i, truth[i]) for i in map(int, r["sample"].split("x"))}
        predicted = called(lookup[r["sample"]], r["method"])
        counts = [len(expected & predicted), len(predicted - expected), len(expected - predicted), len(universe - expected - predicted)]
        if counts != [r[k] for k in ("TP", "FP", "FN", "TN")] or expected != set(r["expected"].split(";")):
            raise ValueError("Independent truth/count audit failed")
    for total in summary:
        rows = [r for r in by_sample if r["view"] == total["view"] and r["method"] == total["method"] and in_group(r, total["group"])]
        if len(rows) != total["samples"] or any(sum(r[k] for r in rows) != total[k] for k in ("TP", "FP", "FN", "TN")):
            raise ValueError("Summary audit failed")
    write_json(root / "COUNT_AUDIT.json", {"truth_parser": "independent stdlib ZIP/XML", "truth_sha256": digest(truth_path),
        "samples": len(calls), "exported_lists_checked": zip_checks, "evidence_rules_replayed": len(evidence),
        "sample_method_view_counts_checked": len(by_sample), "summary_rows_checked": len(summary),
        "fixed_family_universe": len(universe), "all_checks_pass": True})


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--predictions", required=True)
    parser.add_argument("--truth", required=True)
    args = parser.parse_args()
    evaluate(args.predictions, args.truth)
