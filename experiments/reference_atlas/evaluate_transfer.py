"""Local evaluation after all 100 predictions and immutable-input audits finish.

Truth is never uploaded to the prediction host. This is a retrospective method
comparison on the same 100 previously evaluated mixtures, not a new blind test.
"""
import argparse
import csv
import hashlib
import json
import math
from collections import defaultdict
from pathlib import Path


METHODS = ["profile_gene3", "nearest_gene3", "profile_confirmed_baseline",
           "baseline_confirmed", "baseline_z09"]


def measures(tp, fp, fn, tn):
    divide = lambda a, b: a / b if b else None
    denominator = (tp + fp) * (tp + fn) * (tn + fp) * (tn + fn)
    return {"TP": tp, "FP": fp, "FN": fn, "TN": tn,
            "precision": divide(tp, tp + fp), "recall": divide(tp, tp + fn),
            "F1": divide(2 * tp, 2 * tp + fp + fn), "NPV": divide(tn, tn + fn),
            "specificity": divide(tn, tn + fp),
            "MCC": divide(tp * tn - fp * fn, math.sqrt(denominator))}


def write_csv(path, rows):
    with Path(path).open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)


def evaluate(batch, truth_file, output):
    batch, output = Path(batch), Path(output)
    audit = json.loads((batch / "AUDIT.json").read_text())
    assert (batch / "PREDICTIONS_COMPLETE").exists()
    assert audit["all_100_complete"] and audit["all_reused_input_hashes_unchanged"]
    assert audit["code_model_protocol_unchanged"]
    status = json.loads((batch / "STATUS.json").read_text())
    assert len(status) == 100 and all(r["success"] and not r["failed"] for r in status)
    rows = json.loads((batch / "selection.json").read_text())["selected"]
    assert len(rows) == 100
    assert all((batch / "samples" / r["project"] / "SUCCESS").exists() for r in rows)
    # Only after every prediction is complete may the evaluator open truth.
    from openpyxl import load_workbook
    book = load_workbook(truth_file, read_only=True, data_only=True)
    truth = {int(row[0]): str(row[7]).strip() for row in book["Table S1"].iter_rows(
        min_row=4, values_only=True) if isinstance(row[0], int)}
    book.close()
    assert set(truth) == set(range(1, 31))
    truth_hash = hashlib.sha256(Path(truth_file).read_bytes()).hexdigest()
    assert truth_hash == "421cca29377ab01e57cb6aeded7bf446ba4faf06282120a4ed602e8c67bbb01f"
    all_families = set(truth.values()) | {"Cabombaceae", "Capparaceae"}
    calls, panels = {}, {}
    for row in rows:
        path = batch / "samples" / row["project"]
        with (path / "family_evidence.csv").open(newline="", encoding="utf-8") as f:
            panel = {r["family"] for r in csv.DictReader(f)}
        panels[row["sample"]] = panel
        all_families |= panel
        for method in METHODS:
            selected = set((path / (method + ".txt")).read_text().splitlines())
            calls[row["sample"], method] = selected
            all_families |= selected
    first = {}
    by_sample, by_family, coverage = [], [], []
    for row in rows:
        sample = row["sample"]
        composition = tuple(sorted(map(int, sample.split("x"))))
        first.setdefault(composition, sample)
        audit = json.loads((batch / "samples" / row["project"] / "audit.json").read_text())
        coverage.append({"sample": sample, "valid_genes": audit["valid_genes"],
                         "valid_trees": audit["valid_trees"], "parent_orders": len(audit["orders"]),
                         "profile_families": len(panels[sample])})
        for view in ["table_literal", "harmonized_two_taxa"]:
            corrected = {1: "Cabombaceae", 25: "Capparaceae"} if view != "table_literal" else {}
            expected = {corrected.get(i, truth[i]) for i in composition}
            for method in METHODS:
                called = calls[sample, method]
                tp, fp, fn, tn = map(len, (expected & called, called - expected,
                                           expected - called, all_families - expected - called))
                by_sample.append({"sample": sample, "view": view, "method": method,
                    "historical30": row["previously_seen"], "first_composition": first[composition] == sample,
                    **measures(tp, fp, fn, tn), "expected": ";".join(sorted(expected)),
                    "predicted": ";".join(sorted(called)), "false_positives": ";".join(sorted(called - expected)),
                    "false_negatives": ";".join(sorted(expected - called))})
                for family in sorted(all_families):
                    by_family.append({"sample": sample, "view": view, "method": method, "family": family,
                        "expected": family in expected, "selected": family in called,
                        "in_reused_parent_panel": family in panels[sample]})
    summary = []
    for group in ("all100", "unique98", "historical30", "remaining70"):
        for view in ("table_literal", "harmonized_two_taxa"):
            for method in METHODS:
                selected = [r for r in by_sample if r["view"] == view and r["method"] == method and
                    (group == "all100" or group == "unique98" and r["first_composition"] or
                     group == "historical30" and r["historical30"] or group == "remaining70" and not r["historical30"])]
                counts = [sum(r[k] for r in selected) for k in ("TP", "FP", "FN", "TN")]
                summary.append({"group": group, "view": view, "method": method, "samples": len(selected),
                    "negative_universe_families": len(all_families), **measures(*counts),
                    "exact_sets": sum(r["FP"] == r["FN"] == 0 for r in selected)})
    output.mkdir(parents=True, exist_ok=True)
    for name, data in [("summary.csv", summary), ("by_sample.csv", by_sample),
                       ("by_family.csv", by_family), ("coverage.csv", coverage)]:
        write_csv(output / name, data)
    training = json.loads((batch / "model/training.json").read_text())
    protocol = json.loads((batch / "PROTOCOL.json").read_text())
    pct = lambda v: f"{v:.1%}" if v is not None else "NA"
    lines = ["# 100 混样：reference full-profile 科级迭代", "",
        "100/100 预测完成后才在本地打开真表。复用相同目级候选、组装、科级参考、树拓扑和分支长度；未重新跑装配或更改已有预测。",
        "这是同一批已评估样品上的回溯迭代，不能称为新的盲测。100 样品来自相同 30 个来源，包含 98 种组分组合。", "",
        "预先冻结的主方法为 profile_gene3；其他方法保留作对照，不按本轮真值选择主方法。", "",
        "| 口径 | 方法 | TP | FP | FN | Precision | Recall | F1 | NPV | MCC | 完全匹配 |",
        "| --- | --- | ---: | ---: | ---: | ---: | ---: | ---: | ---: | ---: | ---: |"]
    for r in summary:
        if r["group"] == "all100":
            lines.append(f'| {r["view"]} | {r["method"]} | {r["TP"]} | {r["FP"]} | {r["FN"]} | {pct(r["precision"])} | {pct(r["recall"])} | {r["F1"]:.3f} | {pct(r["NPV"])} | {r["MCC"]:.3f} | {r["exact_sets"]}/100 |')
    lines += ["", "原表是主口径；辅助口径只沿用既有 1 号 Cabombaceae、25 号 Capparaceae 的分类统一，不修改原表。",
        f'NPV/TN 的固定全集为本批面板、预测及真表出现过的 {len(all_families)} 个科；该全集有大量容易排除的阴性，NPV 不能单独代表识别能力。', "",
        "## 方法和校准", "",
        "每个基因只选一个确定性哈希最小的现成 exon tree，最多 64 个 reference pseudo-queries；不使用混样标签训练。",
        "pseudo-query 的整个 species identity 从 anchors 移除。训练/校准 query species 按哈希分开；它们仍可能作为其他 pseudo-query 的已知 anchors，属于固定已知参考库上的 transfer 校准，不是前轮严格四分区的新物种评估。",
        "特征包括最近和平均距离、同目竞争间距、去终端偏移的完整 reference 距离指纹残差/相关性/偏移，以及距离排序。family 采用均值而非参考数累加。",
        f'校准阈值为 {training["threshold"]:.6f}，在固定 101 个 cutoffs 上最大化按 species 等权的已知正确接受/整科缺失拒识的平均效用。它是经验拒识阈值，不是物种存在概率或 FDR 保证。',
        f'已知 reference 留出 top-1 准确率 {pct(training["known_top1_accuracy"])}，阈值后的接受比例 {pct(training["known_acceptance"])}。',
        "主方法先用完整距离选每个片段的最近科，再用学习模型验证匹配，要求至少 3 个独立 gene 支持，不按 exon/contig 数累加；nearest_gene3 使用相同 gene 规则但不做模型验证；profile_confirmed_baseline 只保留原最终候选中至少 2 个 gene 获得模型支持者。", "",
        "## 文件与限制", "",
        "summary.csv 含全 100、98 组合去重和历史分组；by_sample.csv / by_family.csv 给出每一处误报、漏检；coverage.csv 报告实际位点、树和父目覆盖。",
        "所有方法使用同一批冻结树，预测后核对所有读取输入的 SHA-256、代码和模型，AUDIT.json 确认一致。",
        "没有调整父目、人工增加真科、按预期组分数截断，或依据本轮结果调阈值。缺失父目、参考覆盖、旁系同源和 assembly 错误仍会限制结果。",
        "模型只改变距离后的评分；本轮不证明建库或测序流水线的端到端提速，也不证明原来二维 pattern 已优于完整距离。",
        "真表 SHA-256：`" + truth_hash + "`。预测 protocol/model 的 SHA-256 与训练说明保存在 batch/PROTOCOL.json 和 batch/model/training.json。"]
    (output / "REPORT.zh-CN.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(json.dumps([r for r in summary if r["group"] == "all100"], indent=2))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--batch", required=True)
    parser.add_argument("--truth", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()
    evaluate(args.batch, args.truth, args.output)
